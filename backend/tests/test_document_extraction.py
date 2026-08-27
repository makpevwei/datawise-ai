import io

import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from app.documents.extraction import (
    ExtractionError,
    extract_docx,
    extract_html,
    extract_pdf,
    extract_pptx,
    extract_text_like,
)
from app.documents.models import DocumentType


def test_extract_txt_preserves_line_metadata():
    content = b"First paragraph line one.\n\nSecond paragraph here.\n"
    segments, warnings = extract_text_like(content, "notes.txt", DocumentType.TXT)
    assert len(segments) == 2
    assert segments[0].location.line_start == 1
    assert warnings == []


def test_extract_md_tracks_headings():
    content = b"# Revenue\nRevenue declined this quarter.\n\n## Regional Detail\nWest region led the decline.\n"
    segments, _ = extract_text_like(content, "report.md", DocumentType.MD)
    headings = [s.location.heading for s in segments]
    assert "Revenue" in headings
    assert "Regional Detail" in headings


def test_extract_text_like_rejects_empty_file():
    with pytest.raises(ExtractionError, match="empty"):
        extract_text_like(b"", "empty.txt", DocumentType.TXT)


def test_extract_text_like_rejects_whitespace_only_content():
    # The latin-1 fallback can decode any byte sequence, so "undecodable"
    # text doesn't exist for this extractor -- but content that decodes to
    # nothing but whitespace/control characters still has no usable text.
    with pytest.raises(ExtractionError, match="empty"):
        extract_text_like(b"\x0c\x0c\x0c", "bad.txt", DocumentType.TXT)


def test_extract_pdf_rejects_non_pdf_bytes():
    with pytest.raises(ExtractionError):
        extract_pdf(b"this is not a pdf", "fake.pdf")


def test_extract_pdf_rejects_empty_bytes():
    with pytest.raises(ExtractionError):
        extract_pdf(b"", "empty.pdf")


def _make_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("Revenue declined 12% in Q2.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "West"
    table.cell(1, 1).text = "39.3M"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx_tracks_headings_and_tables():
    segments, warnings = extract_docx(_make_docx_bytes(), "report.docx")
    assert warnings == []
    headings = [s.location.heading for s in segments]
    assert "Executive Summary" in headings
    assert any("West" in s.text and "39.3M" in s.text for s in segments)


def test_extract_docx_rejects_invalid_file():
    with pytest.raises(ExtractionError):
        extract_docx(b"not a real docx", "fake.docx")


def _make_pptx_bytes() -> bytes:
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Q2 Performance"
    slide.placeholders[1].text = "Revenue declined in the West region."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_extract_pptx_tracks_slide_and_title():
    segments, warnings = extract_pptx(_make_pptx_bytes(), "deck.pptx")
    assert warnings == []
    assert segments[0].location.slide == 1
    assert segments[0].location.slide_title == "Q2 Performance"
    assert "Revenue declined" in segments[0].text


def test_extract_pptx_rejects_invalid_file():
    with pytest.raises(ExtractionError):
        extract_pptx(b"not a real pptx", "fake.pptx")


# -- Python files (Phase 4) ----------------------------------------------------


def test_extract_py_tracks_function_and_class_boundaries():
    content = b'''"""Module docstring."""

def calculate_total(items):
    return sum(items)


class Invoice:
    def __init__(self, amount):
        self.amount = amount
'''
    segments, warnings = extract_text_like(content, "invoice.py", DocumentType.PY)
    assert warnings == []
    headings = [s.location.heading for s in segments]
    assert "def calculate_total" in headings
    assert "class Invoice" in headings


def test_extract_py_keeps_the_def_line_in_the_chunk_body():
    content = b"def compute(x):\n    return x * 2\n"
    segments, _ = extract_text_like(content, "compute.py", DocumentType.PY)
    matching = [s for s in segments if s.location.heading == "def compute"]
    assert matching
    assert "def compute(x):" in matching[0].text


def test_extract_py_rejects_empty_file():
    with pytest.raises(ExtractionError, match="empty"):
        extract_text_like(b"", "empty.py", DocumentType.PY)


# -- HTML files (Phase 4A) -----------------------------------------------------


def test_extract_html_valid_document_with_headings_and_paragraphs():
    html = b"""
    <html><body>
        <h1>Revenue</h1>
        <p>Revenue declined this quarter.</p>
        <h2>Regional Detail</h2>
        <p>West region led the decline.</p>
    </body></html>
    """
    segments, warnings = extract_html(html, "report.html")
    assert warnings == []
    headings = [s.location.heading for s in segments]
    assert "Revenue" in headings
    assert "Regional Detail" in headings
    texts = [s.text for s in segments]
    assert "Revenue declined this quarter." in texts
    assert "West region led the decline." in texts


def test_extract_html_malformed_document_is_still_parsed():
    # Unclosed tags, mismatched nesting -- html.parser is lenient by design.
    html = b"<html><body><h1>Notes<p>Unclosed heading and paragraph<li>loose list item</body>"
    segments, warnings = extract_html(html, "malformed.html")
    assert segments  # extraction did not raise; something usable was recovered
    assert any("loose list item" in s.text for s in segments)


def test_extract_html_preserves_table_content():
    html = b"""
    <html><body>
        <h2>Sales Table</h2>
        <table>
            <tr><th>Region</th><th>Revenue</th></tr>
            <tr><td>North Africa</td><td>18050.80</td></tr>
        </table>
    </body></html>
    """
    segments, _ = extract_html(html, "table.html")
    table_segment = next(s for s in segments if "Region" in s.text)
    assert "North Africa" in table_segment.text
    assert "18050.80" in table_segment.text
    assert table_segment.location.heading == "Sales Table"


def test_extract_html_ignores_script_and_style_content():
    html = b"""
    <html><head><style>.a { color: red; }</style></head>
    <body>
        <script>alert('should never appear in extracted text');</script>
        <p>Real visible paragraph text.</p>
        <!-- a comment that should also never appear -->
    </body></html>
    """
    segments, _ = extract_html(html, "withscript.html")
    all_text = " ".join(s.text for s in segments)
    assert "alert(" not in all_text
    assert "should never appear" not in all_text
    assert "color: red" not in all_text
    assert "a comment that should also never appear" not in all_text
    assert "Real visible paragraph text." in all_text


def test_extract_html_rejects_empty_file():
    with pytest.raises(ExtractionError, match="empty"):
        extract_html(b"", "empty.html")


def test_extract_html_rejects_document_with_no_extractable_text():
    # Only script/style content, nothing extractable once those are stripped.
    html = b"<html><head><style>.a{}</style></head><body><script>1;</script></body></html>"
    with pytest.raises(ExtractionError, match="no extractable text"):
        extract_html(html, "onlyscript.html")


# -- PDF multi-stage extraction (PyMuPDF + OCR) --------------------------------


def test_ocr_capability_detection():
    """OCR capability flags are booleans; they may be True or False depending
    on the environment, but they must never raise."""
    from app.documents.extraction import _ocr_available, _pdf2image_available, _pymupdf_available
    assert isinstance(_pymupdf_available(), bool)
    assert isinstance(_ocr_available(), bool)
    assert isinstance(_pdf2image_available(), bool)


def test_pymupdf_available_in_current_environment():
    """PyMuPDF (pymupdf package) is listed in pyproject.toml and must be
    importable in this test environment."""
    from app.documents.extraction import _pymupdf_available
    assert _pymupdf_available() is True, (
        "PyMuPDF is required and should be installed. "
        "Run: uv add pymupdf"
    )


def test_ocr_available_in_current_environment():
    """pytesseract + tesseract binary must be available in this environment."""
    from app.documents.extraction import _ocr_available
    assert _ocr_available() is True, (
        "pytesseract + tesseract binary must be installed. "
        "Run: uv add pytesseract  and  brew install tesseract"
    )


def _make_minimal_pdf_bytes() -> bytes:
    """Create a minimal valid digital PDF with real selectable text."""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore[no-redef]
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 144), "Revenue declined 12% in Q2 2024.", fontsize=12)
    page.insert_text((72, 164), "West region led the performance gap.", fontsize=12)
    return doc.tobytes()


def test_extract_pdf_digital_text_uses_pymupdf():
    """A digital PDF with native text must be extracted without OCR."""
    pdf_bytes = _make_minimal_pdf_bytes()
    segments, warnings = extract_pdf(pdf_bytes, "digital.pdf")
    assert len(segments) >= 1
    combined = " ".join(s.text for s in segments)
    assert "Revenue" in combined or "revenue" in combined.lower()
    # No OCR warnings expected for a clean digital PDF
    ocr_warnings = [w for w in warnings if "OCR" in w.upper()]
    assert len(ocr_warnings) == 0


def test_extract_pdf_page_numbers_are_preserved():
    """Segment location.page must be populated."""
    pdf_bytes = _make_minimal_pdf_bytes()
    segments, _ = extract_pdf(pdf_bytes, "paged.pdf")
    assert all(s.location.page is not None and s.location.page >= 1 for s in segments)


def test_extract_pdf_raises_on_garbage_bytes():
    """Random bytes must raise ExtractionError, never crash the server."""
    with pytest.raises(ExtractionError):
        extract_pdf(b"\x89PNG garbage not a pdf", "notapdf.pdf")
